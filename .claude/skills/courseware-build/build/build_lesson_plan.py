#!/usr/bin/env python3
"""Generate the Lesson Plan (LP) DOCX for 'Innovating Product Development with 3D
Printing and Additive Manufacturing' (TGS-2022017597) in the Tertiary house format.

Cover page + Document Version Control Record + auto TOC + Arial 11pt body +
colour-coded 2-day schedule tables (9:30am–6:30pm, 8 training hours/day, 1h
lunch, tea within, final assessment Day 2 from 4:00pm). Topics/labs come from
course_data + the domain data files so the LP stays aligned with the deck,
guide and labs.
"""
import os, sys, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
ACT=DOMAIN1+DOMAIN2+DOMAIN3
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
HEADER_FILL="1F6FEB"; TOPIC_FILL="E8F0FE"; BREAK_FILL="FFF4E5"; LUNCH_FILL="FDE9D9"; ASSESS_FILL="E8F7EE"

# slide references from the deck build (slide_map.json is written by build_slides.py)
try:
    with open(os.path.join(HERE,"slide_map.json")) as f: SMAP=json.load(f)
except OSError:
    SMAP={}

def _sl(key):
    n=SMAP.get(key)
    return f" (Slide {n})" if n else ""

def lab_titles(nums):
    return "; ".join(f"Lab {a['num']}: {a['title']}{_sl('lab'+str(a['num']))}"
                     for a in ACT if a['num'] in nums)

def topic_slides(n):
    """Slide range of topic n: from its divider to the slide before the next marker."""
    start=SMAP.get(f"topic{n}")
    if not start: return ""
    nxt=SMAP.get(f"topic{n+1}") or SMAP.get("close") or (SMAP.get("total",0)+1)
    return f" (Slides {start}–{nxt-1})"

# ------------------------------------------------ schedule (single source of truth for timing)
# (start, end, minutes, kind, activity_text)  kind: admin/topic/lab/break/lunch/assess/recap
SCHEDULE = {
 1: (C.DAY_THEMES[1], [
    ("9:30","10:00",30,"admin","Welcome, trainer and learner introductions, course outline, learning outcomes, ground rules and mandatory digital attendance (AM)"),
    ("10:00","11:00",60,"topic","Topic 1 — Overview of Additive Manufacturing: definition, AM vs subtractive, history, the 3D printing workflow, the 7 AM process families, filament/resin/powder printing, DMLS/SLS/SLM (concepts + demo)"+topic_slides(1)),
    ("11:00","11:15",15,"break","Tea break"),
    ("11:15","13:00",105,"lab","Topic 1 — Design for AM in Fusion 360. Hands-on: "+lab_titles([1,2,3])),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:30",90,"lab","Topic 1 — Materials, benefits & trade-offs, case studies and AM selection factors (concepts). Digital attendance (PM). Hands-on: "+lab_titles([4,5])),
    ("15:30","15:45",15,"break","Tea break"),
    ("15:45","17:45",120,"lab","Topic 1 — 3D scanning and reverse engineering: scanning apps, mesh-to-CAD, case studies (concepts). Hands-on: "+lab_titles([6])),
    ("17:45","18:30",45,"topic","Topic 2 begins — Where AM issues come from; the slicer. Hands-on: "+lab_titles([7])+". Day 1 recap and Q&A"),
 ]),
 2: (C.DAY_THEMES[2], [
    ("9:30","9:45",15,"recap","Day 1 recap and mandatory digital attendance (AM)"),
    ("9:45","11:00",75,"topic","Topic 2 — Printer anatomy, bed adhesion, raft/brim/skirt and supports (concepts + demo)"+topic_slides(2)+". Hands-on: "+lab_titles([8])),
    ("11:00","11:15",15,"break","Tea break"),
    ("11:15","13:00",105,"lab","Topic 2 — Extrusion issues, retraction, oozing, temperature, jams, filament types and hybrid metal/non-metal manufacture. Hands-on: "+lab_titles([9,10])),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:30",90,"lab","Topic 3 — Process Improvement: infill, surface quality, textures, supports, post-processing, business impact and sustainability (concepts)"+topic_slides(3)+". Digital attendance (PM). Hands-on: "+lab_titles([11,12,13,14,15,16])),
    ("15:30","15:45",15,"break","Tea break"),
    ("15:45","16:00",15,"assess","Revision, course feedback, TRAQOM survey, Assessment digital attendance and Briefing for Assessment"),
    ("16:00","17:00",60,"assess","Written Assessment (WA) — Short-Answer Questions (SAQ), 1 hour, open book"),
    ("17:00","18:30",90,"assess","Practical Performance (PP) — scenario-based practical tasks, 90 minutes, open book"),
 ]),
}

# ------------------------------------------------ build document
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)

prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("7.0","1 June 2025","Legacy release — v7 trainer deck and schedule.",C.TRAINER),
 ("8.0","18 August 2026",
  "Full regeneration: single-source rebuild aligned to the v8.0 slide deck (150 slides), "
  "16 hands-on labs in per-lab folders, enriched AM content (NIST/ISO-ASTM process families, "
  "materials, case studies) and the WA (SAQ) + PP assessment.",C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "Aligned the 150-slide deck and 16-lab schedule to the supplied diagnostic STL/Fusion/Cura assets; "
  "added evidence-based calibration, support, retraction, cooling, wall, optimisation and finishing tasks.",C.TRAINER),
])
prodoc.add_toc(doc)

def H(text,level=1):
    h=doc.add_heading(text,level=level); return h

H("Course Information",1)
info=[("Course Title",C.TITLE),("WSQ Course Reference",C.COURSE_CODE),
      ("Skills Framework TSC",f"{C.TSC_TITLE} ({C.TSC_CODE})"),
      ("Training Provider",C.ORG+"  ("+C.UEN.replace('UEN: ','UEN ')+")"),
      ("Duration","2 days · 8 training hours per day (16 hours)"),
      ("Daily Timing","9:30 am – 6:30 pm (1-hour lunch; tea breaks within training time)"),
      ("Mode","Instructor-led, hands-on labs with Fusion 360, Ultimaker Cura and FDM 3D printers"),
      ("Trainer",C.TRAINER)]
t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
for k,v in info:
    c=t.add_row().cells; c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
    prodoc._shade_cell(c[0],TOPIC_FILL)
    c[1].text=""; c[1].paragraphs[0].add_run(v).font.size=Pt(10)

H("Learning Outcomes",1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for lo in C.LEARNING_OUTCOMES:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(lo).font.size=Pt(10.5)

H("TSC Abilities and Knowledge",1)
doc.add_paragraph("This course delivers the following TSC abilities and knowledge factors "
                  f"({C.TSC_TITLE}, {C.TSC_CODE}):")
for a in C.TSC_ABILITIES + C.TSC_KNOWLEDGE:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10)

H("Assessment",1)
for a in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],
          "Format: Open Book — course slides, Learner Guide and approved materials only.",
          "Final assessment is conducted on Day 2 from 4:00 pm.",C.ASSESSMENT["note"]]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10.5)

def set_cell(cell,text,bold=False,size=9.5,color=None,fill=None,align=None):
    cell.text=""; p=cell.paragraphs[0]
    if align: p.alignment=align
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name="Arial"
    if color: r.font.color.rgb=color
    if fill: prodoc._shade_cell(cell,fill)

KIND_FILL={"topic":TOPIC_FILL,"break":BREAK_FILL,"lunch":LUNCH_FILL,"assess":ASSESS_FILL,
           "admin":"F3F5F8","recap":"F3F5F8","lab":None}

H("Course Schedule",1)
for day,(theme,rows) in SCHEDULE.items():
    H(f"Day {day} — {theme}",2)
    tbl=doc.add_table(rows=0,cols=3); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=tbl.add_row().cells
    for i,htext in enumerate(["Time","Duration","Topic / Activity"]):
        set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
    tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    training=0
    for start,end,mins,kind,text in rows:
        cells=tbl.add_row().cells; fill=KIND_FILL.get(kind)
        set_cell(cells[0],f"{start}–{end}",bold=(kind in ("topic","assess")),size=9.5,fill=fill)
        set_cell(cells[1],f"{mins} min",size=9.5,fill=fill)
        set_cell(cells[2],text,bold=(kind in ("topic","assess")),size=9.5,fill=fill)
        if kind!="lunch": training+=mins
    for row in tbl.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        row.cells[0].width=Inches(1.15); row.cells[1].width=Inches(0.9); row.cells[2].width=Inches(4.75)
    p=doc.add_paragraph(); r=p.add_run(f"Total training time: {training} minutes ({training//60} hours)."); r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
    assert training==480, f"Day {day} training minutes = {training}, expected 480"

H("Lab Reference (aligned to the TSC)",1)
tt=doc.add_table(rows=0,cols=3); tt.style="Table Grid"
hdr=tt.add_row().cells
for i,htext in enumerate(["Topic","Learning Outcome","Labs"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    cells=tt.add_row().cells
    set_cell(cells[0],f"Topic {tp['code']}: {tp['title']}",bold=True,size=9.5,fill=TOPIC_FILL)
    set_cell(cells[1],tp["weighting"],size=9.5,fill=TOPIC_FILL)
    set_cell(cells[2],", ".join(f"Lab {a['num']}" for a in acts),size=9.5)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
doc.save(OUT)
print("Saved",OUT)
