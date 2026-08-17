#!/usr/bin/env python3
"""Generate the WSQ 'Innovating Product Development with 3D Printing and Additive
Manufacturing' slide deck (all-white Tertiary house style).

Design helpers are the house component set (cover, section, content, two_col,
cards3, big_statement, tile_grid, flow_h, trainer_slide, img_points, img_full,
table_slide, brk). Content is driven by course_data.py + data_domain1..3.py so
the deck stays 100% aligned with the LP, LG and labs. Lab step-by-step detail
lives ONLY in the Learner Guide — the deck carries overview, visuals and the
verify criterion per lab.
"""
import os, sys, math, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3

def _find_repo(start):
    env = os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env):
        return env
    d = start
    for _ in range(8):
        d = os.path.dirname(d)
        if os.path.isdir(os.path.join(d, "courseware")) and os.path.isdir(os.path.join(d, "labs")):
            return d
    return os.path.dirname(os.path.dirname(HERE))
REPO = _find_repo(HERE)
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
LEGACY = os.path.join(REPO, "courseware", "assets", "legacy")

# ---------------- palette ----------------
BLUE=RGBColor(0x1F,0x6F,0xEB); TEAL=RGBColor(0x10,0xB9,0x81); AMBER=RGBColor(0xF5,0x9E,0x0B)
INK=RGBColor(0x16,0x1B,0x26); GREY=RGBColor(0x5B,0x63,0x72); LIGHT=RGBColor(0xF5,0xF8,0xFC)
WHITE=RGBColor(0xFF,0xFF,0xFF); LINE=RGBColor(0xE2,0xE8,0xF0); VIOLET=RGBColor(0x7C,0x3A,0xED)
RED=RGBColor(0xDC,0x26,0x26)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height
BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def rect(s,x,y,w,h,color,line=None):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(1)
    sp.shadow.inherit=False; return sp
def oval(s,x,y,w,h,color):
    sp=s.shapes.add_shape(9,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space=4):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,line in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(space)
        for t,sz,col,bold in line:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold
            r.font.color.rgb=col; r.font.name="Arial"
    return tb
def bullets(s,x,y,w,h,items,size=18,color=INK,gap=10,mcolor=BLUE):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap)
        lvl=it[1] if isinstance(it,tuple) else 0
        text=it[0] if isinstance(it,tuple) else it
        r=p.add_run(); r.text=("•  " if lvl==0 else "–  ")+text
        r.font.size=Pt(size if lvl==0 else size-2); r.font.color.rgb=color if lvl==0 else GREY
        r.font.name="Arial"; r.font.bold=(lvl==0 and isinstance(it,tuple) and len(it)>2 and it[2])
    return tb

PAGE={"n":1}   # cover occupies slide 1
def footer(s):
    PAGE["n"]+=1
    txt(s,Inches(0.4),Inches(7.05),Inches(7.5),Inches(0.35),
        [[(f"{C.TITLE[:60]}  ·  {C.COURSE_CODE}",9,GREY,False)]])
    txt(s,Inches(5.0),Inches(7.05),Inches(3.3),Inches(0.35),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd",9,GREY,False)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(12.4),Inches(7.05),Inches(0.6),Inches(0.35),
        [[(str(PAGE["n"]),9,GREY,False)]],align=PP_ALIGN.RIGHT)

def _fit_title(t):
    n=len(t)
    if n<=48: return 29
    if n<=58: return 25
    if n<=70: return 22
    return 20
def head(s,title,kicker=None,kcolor=BLUE):
    rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),Inches(1.55),kcolor)
    if kicker: txt(s,Inches(0.85),Inches(0.5),Inches(11.6),Inches(0.4),[[(kicker,14,kcolor,True)]])
    txt(s,Inches(0.85),Inches(0.9),Inches(11.9),Inches(0.9),[[(title,_fit_title(title),INK,True)]])
    rect(s,Inches(0.85),Inches(1.7),Inches(11.63),Inches(0.02),LINE)
    return s
def _logo(name):
    p=os.path.join(ASSETS,name)
    return p if os.path.exists(p) else None
def _legacy(name):
    p=os.path.join(LEGACY,name)
    if not os.path.exists(p):
        raise SystemExit(f"missing legacy asset: {name}")
    return p
def _fit(img,max_w,max_h):
    """Return (w,h) in inches fitting the image inside max_w × max_h."""
    with Image.open(img) as im: iw,ih=im.size
    r=min(max_w/iw*ih and max_w/iw, max_h/ih)
    r=min(max_w/iw, max_h/ih)
    return iw*r, ih*r
def add_pic(s,img,x,y,max_w,max_h,center=True,border=True):
    w,h=_fit(img,max_w,max_h)
    px=x+(max_w-w)/2 if center else x
    py=y+(max_h-h)/2 if center else y
    pic=s.shapes.add_picture(img,Inches(px),Inches(py),width=Inches(w))
    if border:
        ln=pic.line; ln.color.rgb=LINE; ln.width=Pt(1)
    return pic

# ---------------- slide templates ----------------
def cover():
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),BLUE); rect(s,0,Inches(7.28),SW,Inches(0.22),TEAL)
    org=_logo("tertiary-infotech-logo.png")
    if org: s.shapes.add_picture(org,Inches(0.85),Inches(0.7),height=Inches(1.05))
    add_pic(s,_legacy("s022_efd51769.png"),9.6,0.55,3.2,1.85,center=True)
    txt(s,Inches(0.9),Inches(2.55),Inches(12),Inches(0.6),[[("COURSE SLIDES  ·  WSQ",16,BLUE,True)]])
    txt(s,Inches(0.9),Inches(3.05),Inches(11.9),Inches(1.9),[[(C.TITLE,33,INK,True)]])
    rect(s,Inches(0.92),Inches(4.75),Inches(2.4),Inches(0.06),TEAL)
    txt(s,Inches(0.9),Inches(5.05),Inches(12),Inches(1.4),
        [[(f"WSQ Course Code: {C.COURSE_CODE}   ·   TSC: {C.TSC_TITLE} ({C.TSC_CODE})",16,GREY,False)],
         [("Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W",14,GREY,False)]],space=6)
    txt(s,Inches(0.9),Inches(6.5),Inches(12),Inches(0.4),[[(f"Version {C.VERSION}  ·  {C.VERSION_DATE}",12,GREY,False)]])
    txt(s,Inches(0.9),Inches(6.85),Inches(12),Inches(0.34),[[("© 2026 Tertiary Infotech Academy Pte Ltd. All rights reserved.  ·  www.tertiarycourses.com.sg",10,GREY,False)]])

def section(kicker,title,n,sub=""):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,BLUE)
    rect(s,Inches(0.85),Inches(2.5),Inches(0.14),Inches(2.0),TEAL)
    txt(s,Inches(1.25),Inches(2.55),Inches(11),Inches(0.6),[[(kicker,18,BLUE,True)]])
    txt(s,Inches(1.25),Inches(3.0),Inches(11.4),Inches(1.6),[[(title,38,INK,True)]])
    if sub: txt(s,Inches(1.27),Inches(4.55),Inches(11),Inches(0.9),[[(sub,15,GREY,False)]])
    txt(s,Inches(10.0),Inches(0.7),Inches(2.8),Inches(1.6),[[(n,72,RGBColor(0xE2,0xE8,0xF0),True)]],align=PP_ALIGN.RIGHT)
    footer(s); return s
def content(title,items,kicker=None,size=19):
    s=head(slide(),title,kicker); bullets(s,Inches(0.85),Inches(1.95),Inches(11.6),Inches(4.9),items,size=size); footer(s); return s
def two_col(title,left,right,kicker=None,lhead="",rhead="",lcol=None,rcol=None):
    s=head(slide(),title,kicker)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(4.7),LIGHT); rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(4.7),LIGHT)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(0.1),lcol or BLUE)
    rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(0.1),rcol or TEAL)
    if lhead: txt(s,Inches(1.1),Inches(2.2),Inches(5.2),Inches(0.4),[[(lhead,16,lcol or BLUE,True)]])
    if rhead: txt(s,Inches(7.2),Inches(2.2),Inches(5.0),Inches(0.4),[[(rhead,16,rcol or TEAL,True)]])
    bullets(s,Inches(1.1),Inches(2.75),Inches(5.2),Inches(3.7),left,size=15)
    bullets(s,Inches(7.2),Inches(2.75),Inches(5.05),Inches(3.7),right,size=15,mcolor=TEAL); footer(s); return s
def cards3(title,cards,kicker):
    s=head(slide(),title,kicker); xs=[Inches(0.85),Inches(5.0),Inches(9.15)]
    for i,c in enumerate(cards[:3]):
        x=xs[i]; col=c[0]
        rect(s,x,Inches(1.95),Inches(3.65),Inches(4.7),LIGHT); rect(s,x,Inches(1.95),Inches(3.65),Inches(0.12),col)
        txt(s,x+Inches(0.25),Inches(2.2),Inches(3.2),Inches(0.6),[[(c[1],19,col,True)]])
        bullets(s,x+Inches(0.25),Inches(2.95),Inches(3.2),Inches(3.4),c[2],size=14,mcolor=col,gap=9)
    footer(s); return s
def big_statement(line1,line2,kicker,color=BLUE):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,color)
    txt(s,Inches(1.1),Inches(2.2),Inches(11),Inches(0.5),[[(kicker,16,color,True)]])
    txt(s,Inches(1.1),Inches(2.8),Inches(11.3),Inches(2.4),[[(line1,36,INK,True)]])
    if line2: txt(s,Inches(1.12),Inches(4.9),Inches(11),Inches(1.2),[[(line2,19,GREY,False)]])
    footer(s); return s
PALETTE=[BLUE,TEAL,VIOLET,AMBER]
def tile_grid(title,items,kicker=None,cols=2,size=15,icons=None,accent=BLUE):
    s=head(slide(),title,kicker,kcolor=accent)
    n=len(items); rows=math.ceil(n/cols)
    X0=Inches(0.85); Y0=Inches(1.95); TOTW=Inches(11.63); AREAH=Inches(4.78)
    gx=Inches(0.3); gy=Inches(0.26)
    cw=int((TOTW-gx*(cols-1))/cols); ch=int((AREAH-gy*(rows-1))/rows)
    bd=Inches(0.6)
    for i,it in enumerate(items):
        r=i//cols; c=i%cols
        x=int(X0+(cw+gx)*c); y=int(Y0+(ch+gy)*r); col=PALETTE[i%len(PALETTE)]
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,Inches(0.1),ch,col)
        oval(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,col)
        ic=icons[i] if icons else str(i+1)
        txt(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,[[(ic,19,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        tx=x+Inches(1.08); tw=cw-Inches(1.32)
        if isinstance(it,tuple):
            txt(s,tx,int(y+Inches(0.14)),tw,int(ch-Inches(0.2)),
                [[(it[0],size+2,INK,True)],[(it[1],size-2,GREY,False)]],anchor=MSO_ANCHOR.MIDDLE,space=3)
        else:
            txt(s,tx,int(y+Inches(0.1)),tw,int(ch-Inches(0.16)),[[(it,size,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def flow_h(title,steps,kicker=None,color=BLUE):
    s=head(slide(),title,kicker,kcolor=color)
    n=len(steps); X0=Inches(0.85); TOTW=Inches(11.63); gap=Inches(0.34)
    cw=int((TOTW-gap*(n-1))/n); y=Inches(2.55); ch=Inches(3.15); bd=Inches(0.82)
    for i,st in enumerate(steps):
        x=int(X0+(cw+gap)*i)
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,cw,Inches(0.1),color)
        oval(s,int(x+cw/2-bd/2),int(y+Inches(0.42)),bd,bd,color)
        txt(s,int(x+cw/2-bd/2),int(y+Inches(0.42)),bd,bd,[[(str(i+1),30,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        txt(s,x+Inches(0.16),int(y+Inches(1.55)),cw-Inches(0.32),int(ch-Inches(1.7)),[[(st,13.5,INK,False)]],align=PP_ALIGN.CENTER)
        if i<n-1:
            txt(s,int(x+cw-Inches(0.04)),int(y+ch/2-Inches(0.3)),int(gap+Inches(0.08)),Inches(0.6),
                [[("▶",15,color,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def trainer_slide(kicker,name,role,rows,initials,accent=BLUE):
    s=head(slide(),"About the Trainer",kicker,kcolor=accent)
    lx=Inches(0.85); lw=Inches(3.65)
    rect(s,lx,Inches(1.95),lw,Inches(4.7),LIGHT); rect(s,lx,Inches(1.95),lw,Inches(0.12),accent)
    bd=Inches(1.7); ax=int(lx+(lw-bd)/2)
    oval(s,ax,Inches(2.5),bd,bd,accent)
    txt(s,ax,Inches(2.5),bd,bd,[[(initials,44,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(s,lx+Inches(0.15),Inches(4.55),lw-Inches(0.3),Inches(0.6),[[(name,21,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,lx+Inches(0.15),Inches(5.2),lw-Inches(0.3),Inches(1.2),[[(role,13,GREY,False)]],align=PP_ALIGN.CENTER)
    rx=Inches(4.9); rw=Inches(7.6); ry=Inches(1.95); rh=Inches(4.7)
    n=len(rows); gy=Inches(0.2); th=int((rh-gy*(n-1))/n)
    for i,(label,val) in enumerate(rows):
        y=int(ry+(th+gy)*i); col=PALETTE[i%len(PALETTE)]
        rect(s,rx,y,rw,th,LIGHT); rect(s,rx,y,Inches(0.1),th,col)
        vruns=[(val,14,INK,False)] if val else [("____________________________________________",13,LINE,False)]
        txt(s,rx+Inches(0.32),y,rw-Inches(0.6),th,
            [[(label.upper(),11,col,True)],vruns],anchor=MSO_ANCHOR.MIDDLE,space=3)
    footer(s); return s
def img_points(title,img,points,kicker=None,accent=BLUE,caption=""):
    """Image aspect-fit left (~7.0 × 4.6), 3–4 takeaway tiles right."""
    s=head(slide(),title,kicker,kcolor=accent)
    add_pic(s,img,0.85,1.95,7.0,4.35 if caption else 4.7)
    if caption:
        rect(s,Inches(0.85),Inches(6.35),Inches(7.0),Inches(0.42),LIGHT)
        txt(s,Inches(1.0),Inches(6.40),Inches(6.7),Inches(0.34),[[(caption,11.5,GREY,False)]])
    rx=8.1; rw=4.4; n=len(points); gy=0.22
    th=(4.7-gy*(n-1))/n
    for i,pt in enumerate(points):
        y=1.95+(th+gy)*i; col=PALETTE[i%len(PALETTE)]
        rect(s,Inches(rx),Inches(y),Inches(rw),Inches(th),LIGHT)
        rect(s,Inches(rx),Inches(y),Inches(0.09),Inches(th),col)
        if isinstance(pt,tuple):
            txt(s,Inches(rx+0.28),Inches(y+0.08),Inches(rw-0.5),Inches(th-0.14),
                [[(pt[0],13.5,col,True)],[(pt[1],12,INK,False)]],space=2)
        else:
            txt(s,Inches(rx+0.28),Inches(y),Inches(rw-0.5),Inches(th),
                [[(pt,12.5,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def img_full(title,img,caption="",kicker=None,accent=BLUE):
    s=head(slide(),title,kicker,kcolor=accent)
    add_pic(s,img,0.85,1.95,11.63,4.3 if caption else 4.72)
    if caption:
        rect(s,Inches(0.85),Inches(6.32),Inches(11.63),Inches(0.46),LIGHT)
        txt(s,Inches(1.05),Inches(6.38),Inches(11.2),Inches(0.36),[[(caption,12.5,GREY,False)]])
    footer(s); return s
def img_pair(title,img1,img2,cap1="",cap2="",kicker=None,accent=TEAL):
    s=head(slide(),title,kicker,kcolor=accent)
    add_pic(s,img1,0.85,1.95,5.65,4.25)
    add_pic(s,img2,6.85,1.95,5.65,4.25)
    if cap1:
        rect(s,Inches(0.85),Inches(6.3),Inches(5.65),Inches(0.44),LIGHT)
        txt(s,Inches(1.0),Inches(6.36),Inches(5.4),Inches(0.34),[[(cap1,11.5,GREY,False)]])
    if cap2:
        rect(s,Inches(6.85),Inches(6.3),Inches(5.65),Inches(0.44),LIGHT)
        txt(s,Inches(7.0),Inches(6.36),Inches(5.4),Inches(0.34),[[(cap2,11.5,GREY,False)]])
    footer(s); return s
def table_slide(title,headers,rows,kicker=None,widths=None,accent=BLUE,note=""):
    s=head(slide(),title,kicker,kcolor=accent)
    X0=0.85; TOTW=11.63; y=1.95
    ncol=len(headers)
    widths=widths or [1.0/ncol]*ncol
    nrows=len(rows)+1
    maxh=4.3 if note else 4.72
    rh=min(0.62, maxh/nrows)
    x=X0
    for j,htxt in enumerate(headers):
        w=TOTW*widths[j]
        rect(s,Inches(x),Inches(y),Inches(w),Inches(rh),accent)
        txt(s,Inches(x+0.14),Inches(y),Inches(w-0.28),Inches(rh),[[(htxt,13,WHITE,True)]],anchor=MSO_ANCHOR.MIDDLE)
        x+=w
    for i,row in enumerate(rows):
        ry=y+rh*(i+1); x=X0
        fill=LIGHT if i%2==0 else WHITE
        for j,cell in enumerate(row):
            w=TOTW*widths[j]
            rect(s,Inches(x),Inches(ry),Inches(w),Inches(rh),fill,line=LINE)
            txt(s,Inches(x+0.14),Inches(ry),Inches(w-0.28),Inches(rh),
                [[(cell,12,INK,j==0)]],anchor=MSO_ANCHOR.MIDDLE)
            x+=w
    if note:
        ny=y+rh*nrows+0.18
        rect(s,Inches(X0),Inches(ny),Inches(TOTW),Inches(0.52),LIGHT)
        txt(s,Inches(X0+0.2),Inches(ny+0.05),Inches(TOTW-0.4),Inches(0.42),[[(note,12.5,GREY,False)]])
    footer(s); return s
def activity_overview(a,kicker):
    s=head(slide(),f"Lab {a['num']} — {a['title']}",kicker,kcolor=TEAL)
    rect(s,Inches(0.85),Inches(1.85),Inches(1.7),Inches(0.5),TEAL)
    txt(s,Inches(0.85),Inches(1.9),Inches(1.7),Inches(0.4),[[(f"LAB {a['num']}",16,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(2.75),Inches(1.88),Inches(9.7),Inches(0.5),[[(a["objective"],12.5,GREY,False)]])
    txt(s,Inches(0.85),Inches(2.6),Inches(11.7),Inches(1.5),[[(a["desc"],19,INK,False)]])
    rect(s,Inches(0.85),Inches(4.25),Inches(11.7),Inches(2.0),LIGHT)
    txt(s,Inches(1.1),Inches(4.45),Inches(11),Inches(0.4),[[("You'll build",14,BLUE,True)]])
    txt(s,Inches(1.1),Inches(4.85),Inches(11),Inches(0.6),[[(a["build"],17,INK,True)]])
    txt(s,Inches(1.1),Inches(5.55),Inches(11.2),Inches(0.6),
        [[("Tools:  ",13,GREY,True),(a["services"],13,GREY,False)],
         [("Step-by-step guide: Learner Guide, Lab ",12,GREY,False),(str(a["num"]),12,GREY,True),
          (f"  ·  labs/lab-{a['num']:02d} folder",12,GREY,False)]],space=4)
    footer(s); return s
def test_slide(a,kicker):
    s=head(slide(),f"Lab {a['num']} — {a['title']}",kicker,TEAL)
    rect(s,Inches(0.85),Inches(2.3),Inches(11.7),Inches(2.6),RGBColor(0xE8,0xF7,0xEE))
    txt(s,Inches(1.2),Inches(2.6),Inches(11),Inches(0.5),[[("✅  Test it",20,RGBColor(0x12,0x7A,0x3E),True)]])
    txt(s,Inches(1.2),Inches(3.3),Inches(11),Inches(1.5),[[(a["test"],17,INK,False)]])
    txt(s,Inches(0.85),Inches(5.3),Inches(11.7),Inches(0.5),
        [[("Full steps and troubleshooting: Learner Guide · Lab ",13,GREY,False),(str(a["num"]),13,GREY,True)]])
    footer(s); return s
def brk(kind,dur,color=AMBER):
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),color); rect(s,0,Inches(7.28),SW,Inches(0.22),color)
    rect(s,Inches(5.4),Inches(2.35),Inches(2.53),Inches(0.1),color)
    txt(s,0,Inches(2.75),SW,Inches(1.2),[[(kind,48,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,0,Inches(4.05),SW,Inches(0.8),[[(dur,22,color,True)]],align=PP_ALIGN.CENTER); PAGE["n"]+=1

def lab_visual(a,kicker):
    """One visual slide per lab, from its legacy imgs."""
    imgs=a.get("img") or []
    if len(imgs)>=2:
        img_pair(f"Lab {a['num']} — What You'll Work With",
                 _legacy(imgs[0]),_legacy(imgs[1]),
                 cap1="Reference — the model / setting you'll produce",
                 cap2="Reference — as prepared in the slicer / tool",
                 kicker=kicker)
    elif imgs:
        img_full(f"Lab {a['num']} — What You'll Work With",_legacy(imgs[0]),kicker=kicker,accent=TEAL)

# ============================================================ BUILD
SLIDE_MAP={}
def mark(key): SLIDE_MAP[key]=len(prs.slides._sldIdLst)+1   # next slide's 1-based number

cover()

# ---------------- ADMIN ----------------
section("COURSE ADMINISTRATION","Welcome & Housekeeping","")
tile_grid("Digital Attendance (Mandatory)",[
 ("AM · PM · Assessment","It is mandatory to take the AM, PM and Assessment digital attendance for WSQ-funded courses."),
 ("SSG QR code","The trainer/administrator displays the digital attendance QR code generated from the SSG portal."),
 ("Scan & submit","Scan the QR code with your mobile phone camera and submit your attendance each session."),
 ("75% minimum","A minimum of 75% attendance is required to be eligible for assessment and funding.")],
 kicker="TRAQOM · SSG DIGITAL ATTENDANCE",cols=2,size=15)
trainer_slide("YOUR TRAINER · GENERAL","Your Trainer","General Trainer template —\nto be completed by the trainer",
 [("Name",""),("Title / Designation",""),("Qualifications",""),
  ("Areas of expertise",""),("Training & industry experience",""),("Contact","")],
 initials="?",accent=GREY)
trainer_slide("YOUR TRAINER",C.TRAINER,"Principal Trainer\nTertiary Infotech Academy Pte Ltd",
 [("Role","Principal Trainer and Founder, Tertiary Infotech Academy Pte Ltd"),
  ("Qualifications","PhD — data science, engineering and emerging-technology educator."),
  ("Delivers","WSQ courses on 3D printing & additive manufacturing, CAD/product design, data science and AI."),
  ("Experience","20+ years of industry training across engineering, manufacturing and technology domains.")],
 initials="AA",accent=BLUE)
content("Let's Know Each Other",[
 "Your name and organisation / role.",
 "Your experience with CAD, 3D printing or manufacturing (if any).",
 "What you want to design, prototype or manufacture after this course."],kicker="ICE-BREAKER")
tile_grid("Ground Rules",[
 "Set your mobile phone to silent mode.","Participate actively — no question is too small.",
 "Mutual respect: agree to disagree.","One conversation at a time.",
 "Be punctual; return from breaks on time.","75% attendance is required."],
 kicker="HOUSEKEEPING",cols=2,size=15)
# Download Course Material — visual browser-mock + steps (LMS screenshot)
img_points("Download Course Material",_legacy("s015_8266d9af.png"),[
 ("1 · Visit the portal","Go to https://lms-tms.tertiaryinfotech.com in your browser."),
 ("2 · Login with OTP","Enter your registered email and key in the One-Time Password sent to you."),
 ("3 · Open your course","Select this course under My Courses."),
 ("4 · Download materials","Download the slides and Learner Guide for the open-book assessment.")],
 kicker="COURSE PORTAL · LMS/TMS",caption="lms-tms.tertiaryinfotech.com — your course portal login page")
two_col("Lesson Plan — 2 Days, 8 Hours/Day",[
 (f"Day 1 — {C.DAY_THEMES[1]}",0),
 ("Digital Attendance (AM)",1),
 ("Trainer and learner introductions · Learning outcomes",1),
 ("Topic 1: Overview of AM and 3D Scanning (Labs 1–6)",1),
 ("Lunch break · Digital Attendance (PM)",1),
 ("Topic 2 begins: Process Related Issues (Lab 7)",1)],
 [(f"Day 2 — {C.DAY_THEMES[2]}",0),
 ("Digital Attendance (AM)",1),
 ("Topic 2: Process Related Issues for AM (Labs 8–10)",1),
 ("Topic 3: Process Improvement for AM (Labs 11–16)",1),
 ("Revision · Course feedback · TRAQOM survey",1),
 ("Digital Attendance (Assessment) · Final Assessment (WA + PP)",1)],
 kicker="SCHEDULE",lhead="Day 1",rhead="Day 2")
tile_grid("Skills Framework — TSC",[
 ("TSC Title",C.TSC_TITLE),
 ("TSC Code",C.TSC_CODE),
 ("Framework","Skills Framework for Precision Engineering — Additive Manufacturing"),
 ("Proficiency","Level 4 — analyse, determine, evaluate and optimise AM processes")],
 kicker="SKILLS FRAMEWORK",cols=2,size=15)
content("TSC Abilities",C.TSC_ABILITIES,kicker="SKILLS FRAMEWORK · ABILITIES",size=16)
content("TSC Knowledge",C.TSC_KNOWLEDGE,kicker="SKILLS FRAMEWORK · KNOWLEDGE",size=16)
tile_grid("Learning Outcomes",[
 ("LO1 — Analyse AM & 3D scanning","Analyze the AM and evaluate the use of reverse engineering using 3D scanning."),
 ("LO2 — Analyse process issues","Analyze the issues and determine the requirements for AM processes."),
 ("LO3 — Optimise & assess","Apply AM optimization techniques and assess the impact of process improvement.")],
 kicker="WHAT YOU'LL ACHIEVE",cols=1,size=16)
cards3("Course Outline",[
 (BLUE,"Topic 1 · Overview",["Overview of AM systems","Material considerations for AM","Benefits and trade-offs of AM","AM with 3D scanning","Design for AM in Fusion 360","Labs 1–6"]),
 (TEAL,"Topic 2 · Process Issues",["Printer hardware & slicing","Bed adhesion: skirt/brim/raft","Extrusion & temperature issues","Filaments and materials","Hybrid metal + polymer parts","Labs 7–10"]),
 (VIOLET,"Topic 3 · Improvement",["Infill and wall optimisation","Surface quality & textures","Support strategy","Post-processing chains","Business impact & sustainability","Labs 11–16"])],kicker="2 DAYS · 16 HOURS")
tile_grid("Briefing for Assessment",[
 "Place phones and other materials under the table or on the floor.",
 "No photos or recording of assessment scripts.","No discussion during the assessment.",
 "Use a black/blue pen for hard-copy assessments.","No liquid paper / correction tape.",
 "Scripts are collected when time is up."],
 kicker="ASSESSMENT RULES",cols=2,size=14)
tile_grid("Assessment",[
 ("Written Assessment (WA)","Short-Answer Questions (SAQ) — 1 hour, open book."),
 ("Practical Performance (PP)","Scenario-based practical tasks — 90 minutes, open book."),
 ("Open book","Slides, Learner Guide and approved materials only."),
 ("Eligibility","Minimum 75% attendance to be eligible for assessment and funding."),
 ("Appeal","An appeal process is available if required."),
 ("Outcome","Competent (C) / Not Yet Competent (NYC) recorded on the Assessment Summary Record.")],
 kicker="FINAL ASSESSMENT",cols=2,size=13)
flow_h("Assessment Flow",[
 "TRAQOM survey — scan the QR code on the LMS",
 "Assessment digital attendance — scan the SSG QR",
 "Sit WA (SAQ) then the Practical Performance — open book",
 "Upload your answers on the LMS",
 "Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY")

# ================= TOPIC 1 =================
T=C.TOPICS[0]
mark("topic1")
section("TOPIC 01",T["title"],"01",T["subtitle"])
tile_grid("What is Additive Manufacturing?",[
 ("Layer-by-layer fabrication","AM builds parts by depositing material layer upon layer, directly from a computer-aided design (CAD) file."),
 ("No part-specific tooling","No dies, molds or fixtures — the digital file is the tooling, so design changes cost nothing to retool."),
 ("Freeform geometry","Multiple layers built in the X–Y plane stack along Z, enabling shapes machining cannot reach."),
 ("Rapid & flexible","Concept-to-part in hours enables fast prototyping, iteration and mass customisation.")],
 kicker="OVERVIEW · K1",cols=2,size=15)
two_col("Additive vs Subtractive Manufacturing",[
 ("Additive (AM)",0),("Builds the part up layer by layer",1),
 ("Material only where needed — minimal scrap",1),
 ("Complexity is essentially free",1),
 ("No tooling — economic at lot size one",1),
 ("Slower per part at high volumes",1)],
 [("Subtractive (CNC/machining)",0),("Cuts material away from a solid block",1),
 ("Significant chip/scrap waste",1),
 ("Complex internal features often impossible",1),
 ("Tooling & setup favour large volumes",1),
 ("Excellent finish and tight tolerances",1)],
 kicker="TWO PHILOSOPHIES · K1",lhead="Add material",rhead="Remove material")
tile_grid("3D Printing vs Additive Manufacturing",[
 ("3D printing","Strictly, the build step — the machine depositing layers. The term the consumer world uses."),
 ("Additive manufacturing","The full industrial workflow: design, simulation, build, post-processing, inspection and certification."),
 ("Same core, different scope","Every AM part is 3D printed — but AM adds the engineering discipline around the print."),
 ("In this course","We use both terms; Topic 2–3 cover the process discipline that turns printing into manufacturing.")],
 kicker="TERMINOLOGY",cols=2,size=15)
flow_h("A Brief History of 3D Printing",[
 "1981 — Hideo Kodama publishes the first layer-by-layer photopolymer method",
 "1984–87 — Chuck Hull invents stereolithography (SLA) and ships the first commercial printer",
 "1988–92 — FDM patented by Scott Crump; SLS emerges for powders",
 "2009 — key FDM patents expire; RepRap & desktop printing explode",
 "Today — metal AM flies in aircraft, implants and rockets"],
 kicker="MILESTONES",color=VIOLET)
flow_h("The 3D Printing Workflow",[
 "DESIGN — model the part in CAD (Fusion 360) or capture it by 3D scanning",
 "EXPORT — tessellate to an STL/3MF mesh file",
 "SLICE — the slicer (Cura) converts the mesh to layer toolpaths (G-code)",
 "PRINT — the machine builds the part layer by layer",
 "POST-PROCESS — remove supports, cure, finish, inspect"],
 kicker="CAD TO PART",color=BLUE)
tile_grid("The 7 AM Process Families (ISO/ASTM 52900)",[
 ("Material Extrusion","FDM/FFF — molten filament laid down in beads. Desktop standard."),
 ("Vat Photopolymerisation","SLA, DLP, LCD/MSLA — light cures liquid resin, layer by layer."),
 ("Powder Bed Fusion","SLS, DMLS, SLM, EBM — laser/e-beam fuses powder layers."),
 ("Material Jetting","PolyJet/NanoParticle jetting — droplets of material UV-cured in place."),
 ("Binder Jetting","Liquid binder glues powder — metals, sand moulds, full-colour parts."),
 ("Directed Energy Deposition","Blown powder or wire melted by laser/arc — repair & large metal parts."),
 ("Sheet Lamination","LOM, ultrasonic consolidation — sheets bonded and cut."),
 ("Hybrid systems","AM + CNC finishing in one machine — additive build, subtractive precision.")],
 kicker="TECHNOLOGY MAP · K1",cols=2,size=13)
tile_grid("Vat Photopolymerisation & Powder Bed Fusion",[
 ("SLA","A UV laser cures resin one small spot at a time — fine detail, smooth finish."),
 ("DLP / LCD (MSLA)","A projector or LCD mask cures a whole layer at once — speed at pixel resolution."),
 ("SLS","A laser sinters polymer powder — tough nylon parts, no supports needed."),
 ("LPBF / DMLS / SLM","Laser powder bed fusion of metals — sintered (DMLS) or fully melted (SLM)."),
 ("EBM","An electron beam melts metal powder in vacuum — aerospace titanium."),
 ("MJF","HP Multi Jet Fusion — agent + IR lamps fuse nylon powder fast.")],
 kicker="PROCESS FAMILIES 2–3 · K1",cols=2,size=13)
img_points("Printing with Filament (FDM)",_legacy("s022_efd51769.png"),[
 ("Layer on layer","Each layer prints on top of the previous one, up from the build platform — like building a wall."),
 ("Needs a base","The first layer must stick to the platform; poor bases need adhesion aids or redesign."),
 ("Overhangs need support","Sudden overhangs print into mid-air — the slicer generates removable support structures."),
 ("Design to avoid supports","Orient and design parts to minimise supports: faster prints, better surfaces.")],
 kicker="TOPIC 01 · FDM",caption="Desktop FDM printers — cartesian and delta architectures")
img_points("Printing with Resin (SLA / DLP / LCD)",_legacy("s025_c2c7fe92.jpg"),[
 ("UV-curable resin","A tray of liquid photopolymer is selectively hardened by a light source."),
 ("Laser vs projector","SLA cures spot by spot; DLP/LCD cure an entire layer at once — features as fine as a pixel."),
 ("Printed solid","Resin parts usually print solid or drained — trapped liquid resin cannot stay inside."),
 ("Post-cure required","Wash in IPA, then UV post-cure to reach full strength — always with gloves.")],
 kicker="TOPIC 01 · RESIN",caption="A desktop SLA resin printer with its UV-blocking cover")
img_points("Printing with Powder (SLS / Metal PBF)",_legacy("s027_ad621248.png"),[
 ("Most aggressive complexity","Unfused powder supports the part, freeing the most complex geometries in AM."),
 ("Industrial territory","Powder is messy and metal powders are an explosion hazard — not a casual in-house setup."),
 ("Post-processing heavy","Depowdering, heat treatment and finishing depend on technology and material."),
 ("Nylon to titanium","From SLS nylon ducting to DMLS titanium implants and rocket injectors.")],
 kicker="TOPIC 01 · POWDER",caption="An SLS part emerging from its unsintered powder cake")
cards3("Metal Powder Bed Fusion — Three Siblings",[
 (BLUE,"DMLS — sinter",["Laser fuses metal powder by sintering — heat without full melting","Stainless, cobalt-chrome, aluminium, Inconel","Dental implants, aerospace brackets"]),
 (TEAL,"SLS — polymers",["Same principle applied to thermoplastic powders","Tough nylon (PA12) parts","No support structures — powder is the support"]),
 (VIOLET,"SLM — full melt",["Laser fully melts the metal to a liquid state","Denser parts, higher residual stress","Demands tight process control"])],
 kicker="DMLS · SLS · SLM · K1/K2")
tile_grid("AM Materials Palette",[
 ("Polymers","PLA, ABS, PETG, Nylon, TPU filaments; photopolymer resins; PA12 powders."),
 ("Metals","Stainless & tool steels, titanium, aluminium, cobalt-chrome, Inconel, copper."),
 ("Ceramics & composites","Technical ceramics, carbon/glass-fibre filled polymers, metal-matrix blends."),
 ("Specialty","Wood/metal/stone-filled, conductive, magnetic, glow, flexible TPE, bio-inks.")],
 kicker="MATERIALS · K2/K3",cols=2,size=14)
tile_grid("Selecting an AM Solution — Decision Factors",[
 ("Material properties","Mechanical strength, temperature, chemical resistance and biocompatibility the part must deliver."),
 ("Design complexity","Feature detail, internal channels, overhangs — some processes resolve far finer detail than others."),
 ("Volume & speed","Prototyping favours AM at lot size one; beyond ~100s of units, moulding usually wins."),
 ("Post-processing","Supports, curing, machining, finishing — each process carries its own downstream cost."),
 ("Cost structure","Machine, material, labour and energy per part — plus capital outlay and maintenance."),
 ("Learning curve","Operator skill, software toolchain and process maturity in your team.")],
 kicker="K3 · CHOOSING WISELY",cols=2,size=13)
img_points("Considerations When Buying a 3D Printer",_legacy("s032_45f34e07.png"),[
 ("Cost","Machine price, materials, spares and running cost per printed hour."),
 ("Learning curve","From plug-and-print desktop FDM to industrial PBF needing trained operators."),
 ("Build volume","Big printers cost and occupy exponentially more — match to your part sizes."),
 ("Capital & maintenance","Ventilation, power, calibration schedules and service contracts.")],
 kicker="K3 · PROCUREMENT",caption="A desktop FDM printer mid-print")
tile_grid("Benefits of Additive Manufacturing",[
 ("Design freedom","Complex geometry is free — internal lattices, conformal channels, organic forms."),
 ("Rapid prototyping","Idea to physical part in hours; iterate daily instead of quarterly."),
 ("Mass customisation","Every part can differ at no extra cost — dental, hearing aids, prosthetics."),
 ("Less waste","Material is placed only where needed — high buy-to-fly ratios beaten dramatically."),
 ("No tooling","Skip molds and dies entirely — economic at lot size one, instant design changes."),
 ("Distributed production","Ship files, not parts — print at the point of need.")],
 kicker="K6 · WHY AM WINS",cols=2,size=13)
tile_grid("Trade-offs of Additive Manufacturing",[
 ("Throughput","Layer-by-layer is slow — injection moulding beats AM beyond hundreds of units."),
 ("Mechanical anisotropy","Parts are typically 10–50% weaker across layer boundaries than moulded equivalents."),
 ("Material range","Far fewer qualified materials than conventional manufacturing, and often costlier."),
 ("Surface & accuracy","Layer lines and dimensional tolerance often demand post-processing."),
 ("Post-processing load","Support removal, curing, heat treatment — hidden time and labour."),
 ("Repeatability & QA","Process variation makes certification harder — a live research frontier (NIST).")],
 kicker="K6 · EYES OPEN",cols=2,size=13)
img_pair("Case Study — Digital Dentistry",_legacy("s033_bc1267e3.png"),_legacy("s034_0154292f.jpg"),
 cap1="3D-printed dental models and appliances",cap2="Scan-to-model workflow: intraoral scan → print",
 kicker="CASE STUDY · K4/K6",accent=VIOLET)
img_pair("Case Study — Hearing Aids",_legacy("s035_f7ec68a3.png"),_legacy("s036_bb7b7064.jpg"),
 cap1="Ear impression captured for a custom fit",cap2="Custom-fit 3D-printed ear tips",
 kicker="CASE STUDY · MASS CUSTOMISATION · 99% OF SHELLS ARE 3D-PRINTED",accent=VIOLET)
img_pair("Case Study — Prosthetics for Humans and Animals",_legacy("s037_9ec5379c.png"),_legacy("s038_29778e68.jpg"),
 cap1="Printed prosthetic hand components",cap2="A dog fitted with 3D-printed prosthetic legs",
 kicker="CASE STUDY · ACCESSIBLE ENGINEERING",accent=VIOLET)
img_points("Case Study — Distributed Manufacturing",_legacy("s039_5e8797ce.png"),[
 ("Ship files, not parts","Physical goods distributed as downloadable files — like music and movies."),
 ("Print at point of need","Spare parts printed on-site cut inventory, logistics and lead time."),
 ("Real constraints","Quality and materials remain limited; printers need calibration and maintenance."),
 ("People still matter","Staff must be trained, filament supplied, tech support provided.")],
 kicker="CASE STUDY · K4",caption="A replacement part printed where it is needed")
tile_grid("AM Across Industries",[
 ("Aerospace","Lightweight blisks, brackets and fuel nozzles — mass down, lead time down."),
 ("Automotive","Rapid prototypes, custom jigs & fixtures, low-volume spares."),
 ("Medical & dental","Patient-specific implants, surgical guides, aligners, hearing aids."),
 ("Rail & heavy industry","Bespoke obsolete spares and worn-part repair by DED."),
 ("Construction","Extruded-concrete walls and printed architectural elements."),
 ("Consumer & education","Desktop prototyping, maker culture, STEM classrooms.")],
 kicker="K4 · WHO USES IT",cols=2,size=13)
big_statement("Complexity is free. Customisation is free. Volume is not.",
 "AM rewrites the economics of scope and scale — prototype and personalise with AM; mass-produce with molds.","THE DELOITTE LENS · K4/K6",color=BLUE)
img_points("3D Scanning — Capture Reality",_legacy("s042_23187125.png"),[
 ("Photogrammetry","Dozens of photos from all angles reconstructed into a textured 3D mesh."),
 ("LiDAR / structured light","Depth sensors on phones and dedicated scanners capture geometry directly."),
 ("Mobile apps","Polycam, Kiri Engine, Scandy Pro, Heges, ItSeez3D, EM3D, Trnio, Capture 3D."),
 ("Output","A mesh (STL/OBJ) ready for repair, remodelling and re-printing.")],
 kicker="3D SCANNING · A4",caption="Photogrammetry: phone capture → reconstructed bust")
tile_grid("3D Model Software Landscape",[
 ("TinkerCAD","Free browser CAD — drag-and-drop solids, perfect first step."),
 ("Fusion 360","This course's tool — parametric sketch-driven CAD/CAM with free personal licence."),
 ("Inventor / SolidWorks","Professional parametric MCAD for production engineering."),
 ("Onshape","Full cloud CAD — CAD in a browser tab, built-in PDM."),
 ("SketchUp","Fast push-pull architectural modelling."),
 ("GrabCAD / Thingiverse / YouMagine","Model repositories — millions of ready-to-print designs.")],
 kicker="TOOLBOX",cols=2,size=13)
img_points("Fusion 360 — Sketch-Driven CAD",_legacy("s044_8e91017a.jpg"),[
 ("Sketch first","2D shapes drawn on sketch planes become 3D parts through modelling operations."),
 ("Dimension-accurate","Parametric dimensions drive printable, engineering-accurate geometry."),
 ("Four core techniques","Extrude, Revolve, Sweep and Loft cover most printable parts — Labs 1–4."),
 ("Beyond solids","Freeform (T-spline) and surface modelling handle organic, complex curvature.")],
 kicker="DESIGN FOR AM · A1",caption="Fusion 360 modelling a consumer product enclosure")
tile_grid("Fusion 360 — The Four Core Techniques",[
 ("Extrude — Lab 1","Draw a 2D profile, push it along the third axis into a solid — the bread-and-butter operation."),
 ("Sweep — Lab 2","Move a profile along a 2D path on a perpendicular plane — tubes, clips, wires."),
 ("Revolve — Lab 4","Spin a half cross-section around an axis — funnels, vases, pulleys, bottles."),
 ("Loft — Lab 3","Blend between profiles on different planes, guided by rails — freeform vessels and transitions.")],
 kicker="A1 · TECHNIQUES → LABS",cols=2,size=14)
content("Freeform and Surface Modelling",[
 "Freeform (T-spline) modelling shapes organic forms by moving, rotating and scaling the points, edges and faces of primitive bodies.",
 "Surface modelling constructs complex curvature patch by patch, then thickens or stitches into printable solids.",
 "Use them when sketch-based solids cannot reach the shape — ergonomic grips, sculpted housings, scanned-form matching.",
 "Lab 6's scanned meshes and the torso case study are edited with exactly these tools."],
 kicker="BEYOND SKETCHES · A1",size=16)
tile_grid("Designs with Stress Concentration",[
 ("What it is","A point where local stress runs far higher than the surrounding material."),
 ("Where it hides","Sharp corners, holes, notches and grooves — any sudden change in geometry."),
 ("Why AM cares","Layer boundaries + stress risers = the crack starts exactly there."),
 ("The fix","Fillet radii, tapers and relief notches ease the transition and spread the load.")],
 kicker="DESIGN QUALITY · A1",cols=2,size=15)
img_points("Design Tips to Minimise Stress Concentration",_legacy("s055_27a4e58b.png"),[
 ("Fillet every corner","Use fillet radii or tapers to ease transitions between shapes."),
 ("Relief notches","A round relief notch stops a crack from starting in a sharp bend."),
 ("Avoid sharp corners","Especially at load-bearing intersections and layer boundaries."),
 ("Orient for strength","Put tensile loads along the layers, not across them.")],
 kicker="DESIGN QUALITY · A1",caption="A relief notch in a bend corner de-concentrates stress")
img_points("Digital and Physical Tolerance",_legacy("s056_1202dfea.png"),[
 ("Design the gap","Assembling parts need designed clearance — 0.2 mm is the FDM starting point."),
 ("Measure digitally","Fusion 360's Measure, Section Analysis and sketch tools verify the designed gap."),
 ("Verify physically","Digital calipers confirm the printed part reproduced the digital tolerance."),
 ("Lab 5","You'll design, print and measure a 0.2 mm container-and-lid fit.")],
 kicker="FIT & ASSEMBLY · A1",caption="Wireframe visual display of a lid-to-box clearance measurement")
flow_h("Reverse Engineering with 3D Scanning",[
 "IDENTIFY — inspect and repair damage on the original part",
 "SCAN — capture it with a high-resolution 3D scanner",
 "REFINE — clean the mesh into an accurate CAD model",
 "SELECT — pick the AM process for the material & design",
 "MANUFACTURE — print and post-process the component"],
 kicker="A4 · THE WORKFLOW",color=VIOLET)
img_points("Case Study — Reverse Engineering a Mesh",_legacy("s060_84bbd3d7.png"),[
 ("Import the scan","A ready-made mesh/scan imports straight into Fusion 360."),
 ("Section the mesh","The mesh section tool produces a clean cross-section curve."),
 ("Trace & rebuild","Trace the section with a corrected sketch — fix wear and damage digitally."),
 ("Mirror & revolve","Rebuild the part as clean parametric CAD derived from the scan.")],
 kicker="A4 · MESH → CAD",caption="Sectioning a scanned piston mesh to recover its true profile")
img_points("Case Study — Reverse Engineering an Organic Form",_legacy("s061_78548b45.png"),[
 ("Freeform cylinder","Start a freeform cylinder over the torso scan and match its contours."),
 ("Pull to the scan","The Pull tool snaps the cylinder's vertices onto the scanned surface."),
 ("Thicken","Thicken the matched surface into a printable, wearable shell."),
 ("Result","Custom-fit body armour conformal to the scanned torso.")],
 kicker="A4 · ORGANIC FORMS",caption="T-spline body matched to a torso scan, ready to thicken")
img_points("3D Printing Repositories",_legacy("s062_dc2d5916.png"),[
 ("Thingiverse","thingiverse.com — the largest free STL library."),
 ("YouMagine","youmagine.com — Ultimaker's open design community."),
 ("Printables / GrabCAD","Curated models and engineering-grade CAD downloads."),
 ("Check the licence","Respect Creative Commons terms before commercial use.")],
 kicker="READY-MADE MODELS",caption="Browse, download, slice — millions of proven printable designs")

# ---- Topic 1 labs
cards3("Hands-On Labs — Topic 1",[
 (BLUE,"Design for AM",["Lab 1: Extrude — keychain token","Lab 2: Sweep — paper clip","Lab 3: Loft — plant pot"]),
 (TEAL,"Design precision",["Lab 4: Revolve — funnel","Lab 5: Tolerance — container & lid"]),
 (VIOLET,"Capture reality",["Lab 6: 3D scan a real object","Export, repair and print the scan"])],
 kicker="WHAT YOU'LL DO · LABS 1–6")
for a in DOMAIN1:
    kick=f"TOPIC 01 · HANDS-ON"
    mark(f"lab{a['num']}")
    activity_overview(a,kick)
    lab_visual(a,kick)
    test_slide(a,f"LAB {a['num']} · VERIFY")
content("Recap — Topic 1",[
 "AM builds parts layer by layer from CAD with no part-specific tooling — seven process families cover polymers to metals.",
 "Material, complexity, volume, post-processing and cost drive the selection of an AM solution.",
 "AM's benefits (freedom, speed, customisation, low waste) trade against throughput, anisotropy and finishing.",
 "3D scanning + Fusion 360 turn real objects into editable CAD for re-manufacture — reverse engineering.",
 "You modelled with extrude, sweep, loft and revolve, engineered a 0.2 mm fit, and printed a scanned object."],
 kicker="TOPIC 01 · RECAP",size=16)

# ================= TOPIC 2 =================
T=C.TOPICS[1]
mark("topic2")
section("TOPIC 02",T["title"],"02",T["subtitle"])
tile_grid("Where AM Issues Come From",[
 ("Printer hardware","Frame, motion system, extruder, hotend, bed — wear and misalignment show up in every print."),
 ("Materials","Diameter variation, moisture, contamination and wrong temperature windows."),
 ("Printing process conditions","Speed, temperature, retraction, cooling and adhesion settings."),
 ("Environment conditions","Draughts, ambient temperature, humidity and vibration."),
 ("Design-caused","Overhangs, thin walls, sharp corners and tolerance choices that fight the process.")],
 kicker="FIVE SOURCES · K2",cols=2,size=13)
img_full("3D Printer Anatomy",_legacy("s066_7dc22cec.jpg"),
 caption="Extruder · hotend · X/Y/Z steppers · build plate · power supply · mainboard · controller — know the part to diagnose the fault",
 kicker="HARDWARE · K2")
img_points("The Slicer — From Mesh to G-code",_legacy("s067_aab1aa3b.png"),[
 ("What it does","Converts the 3D model (STL) into specific layer-by-layer instructions for the printer."),
 ("G-code","The output is G-code — movement, temperature and extrusion commands the printer executes."),
 ("Settings = quality","Every quality, speed and strength decision is made here, before the printer moves."),
 ("Cura","This course uses Ultimaker Cura — free, powerful, industry-standard for FDM.")],
 kicker="SOFTWARE · K2",caption="Cura preparing a hull model — layer preview of the sliced toolpath. Alternatives: PrusaSlicer, SelfCAD, OctoPrint, IceSL, KISSlicer")
img_points("Getting the Print to Stick",_legacy("s077_5df8a06c.png"),[
 ("Print geometry","Small or complex bases have little contact area to grip the bed."),
 ("Nozzle distance","First-layer gap must be right — too far and nothing sticks, too close and it drags."),
 ("Platform surface","Bed material and films/coatings must be compatible with your filament."),
 ("Platform temperature","A heated bed keeps the first layers relaxed and attached.")],
 kicker="FIRST LAYER · K2/A2",caption="An adhesion failure mid-print — the part broke loose and dragged")
cards3("Raft, Brim and Skirt",[
 (BLUE,"Raft",["A printed grid platform the model sits on","Rescues warping-prone and top-heavy parts","Harder to remove; roughens the bottom face"]),
 (TEAL,"Brim",["A flat border fused to the part's first-layer edge","Boosts adhesion and fights warping","Peels away; light cleanup at the edge"]),
 (VIOLET,"Skirt",["Loops around the part, never touching it","Primes the nozzle, evens the flow","Doubles as a bed-levelling check"])],
 kicker="ADHESION AIDS · A2/A3 · LAB 8")
img_pair("Extrusion Issues — Over vs Under",_legacy("s078_2ed162e1.png"),_legacy("s079_454373d7.png"),
 cap1="Over-extrusion: blobby surfaces, oversized walls, ridged tops",
 cap2="Under-extrusion: gaps, thin walls, missing layers",
 kicker="DIAGNOSIS · A2",accent=RED)
two_col("Diagnosing Extrusion Faults",[
 ("Over-extrusion — too much plastic",0),
 ("Blobby prints; holes come out undersized",1),
 ("Ridges stick up from top surfaces",1),
 ("Fix: calibrate flow %, check filament diameter",1),
 ("Fix: lower temperature a step",1)],
 [("Under-extrusion — too little plastic",0),
 ("Gaps between lines; weak, brittle layers",1),
 ("Extruder clicking or grinding the filament",1),
 ("Fix: check for partial clogs and jam points",1),
 ("Fix: raise temperature / slow down / dry filament",1)],
 kicker="A2/A3 · CAUSE → FIX",lhead="Too much",rhead="Too little",lcol=RED,rcol=BLUE)
img_points("Retraction — Stopping the Flow",_legacy("s082_a1305825.png"),[
 ("Why retract","Plastic in the nozzle is under pressure; jumps across open space would leak."),
 ("The move","The extruder pulls filament back a few millimetres to equalise pressure."),
 ("Distance & speed","Tuned per printer/material — Bowden setups need more than direct drive."),
 ("Get it wrong","Too little → oozing and strings; too much → dimples and jams.")],
 kicker="A2 · RETRACTION",caption="Cura retraction settings — distance and speed")
img_pair("Oozing and Stringing",_legacy("s081_3d5a76a5.png"),_legacy("s084_5a2ace54.png"),
 cap1="Blobs and dimples where pressure equalised badly",
 cap2="Strings bridge the travel moves between towers",
 kicker="A2 · SYMPTOMS",accent=RED)
img_points("Temperature — The Master Variable",_legacy("s088_877f1ce9.png"),[
 ("Too low — way","Extrusion fails: clicking extruder, ground filament, nothing flows."),
 ("Too low — slightly","Layers don't bond; the part delaminates along layer lines."),
 ("Too high — slightly","Bridges and overhangs sag; stringing hard to control."),
 ("Too high — way","Smoking, discolouration, dripping liquid instead of controlled extrusion.")],
 kicker="A2 · TEMPERATURE",caption="Setting printing temperature in Cura's material section")
img_pair("Reading Temperature Faults",_legacy("s085_1b07abaf.png"),_legacy("s089_1490e1c2.png"),
 cap1="Bonding too cold — layers separate under light force",
 cap2="Far too hot — sagging, melted detail, discolouration",
 kicker="A2 · EVIDENCE",accent=RED)
img_points("Filament Jams and Heat Creep",_legacy("s090_cea5cb08.png"),[
 ("The jam","Nozzle moves, nothing extrudes — clicking as the motor skips backwards."),
 ("Heat creep","Heat spreads up the filament path; filament swells and binds before the melt zone."),
 ("PLA is fragile here","PLA's low glass point makes it the most heat-creep-prone material."),
 ("Prevention","Cooling fan on the heatbreak, retraction not overtuned, quality filament.")],
 kicker="A2 · JAMS",caption="Heat creep: the melt zone (red) spreading up the cold side")
img_points("Clearing a Clogged Nozzle",_legacy("s091_0485b0d9.png"),[
 ("Hot needle","Heat the nozzle, insert a needle/wire held with pliers — never fingers."),
 ("Gentle pressure","Poke the clog while pressing the filament lightly from above."),
 ("Cold pull","Once flowing, cool slightly and yank the filament to extract residue."),
 ("Persistent clogs","Swap the nozzle — they are consumables, not heirlooms.")],
 kicker="A3 · PROCEDURE",caption="Clearing a clog with a needle — hot end, pliers, patience")
tile_grid("The Filament Universe",[
 ("Standards","PLA · ABS · PETG/PET · Nylon — the four you must know (next slides)."),
 ("Flexibles","TPE / TPU — rubber-like parts, gaskets, tyres."),
 ("Composites","Carbon-fibre, wood, sandstone, metal-filled — stiffness or looks."),
 ("Soluble","PVA (water) and HIPS (limonene) — dissolvable support materials."),
 ("Functional","Conductive, magnetic iron, glow-in-the-dark, Amphora."),
 ("Engineering","PC, PEEK, ULTEM — high temperature, high strength, high difficulty.")],
 kicker="MATERIALS · K2",cols=2,size=13)
two_col("ABS vs PLA",[
 ("ABS — engineering workhorse",0),
 ("Durable, strong, lightweight and slightly flexible",1),
 ("Cheap and proven — the professional's default for years",1),
 ("Petroleum-based, not biodegradable",1),
 ("High temperatures; warps without heated bed/enclosure",1),
 ("Unpleasant fumes — ventilate the room",1)],
 [("PLA — the friendly starter",0),
 ("Plant-derived; sweet smell, no harsh fumes",1),
 ("Easy to print, minimal warping — ideal for beginners",1),
 ("Huge colour/effect range: glow, translucent, silk",1),
 ("Brittle vs ABS; softens in a hot car",1),
 ("Attracts moisture — store dry",1)],
 kicker="K2 · MATERIALS HEAD-TO-HEAD",lhead="ABS",rhead="PLA")
two_col("PETG vs Nylon",[
 ("PETG — tough all-rounder",0),
 ("Strong, flexible, water- and chemical-resistant",1),
 ("Low warp, low shrink; food-safe grades exist",1),
 ("Doesn't absorb moisture from the air readily",1),
 ("Needs fine-tuned nozzle & bed temperatures",1),
 ("Strings more than PLA while tuning",1)],
 [("Nylon — engineering flexibility",0),
 ("Strong, durable and flexible; less brittle than ABS/PLA",1),
 ("100% thermoplastic — re-meltable without bond loss",1),
 ("High melting temperature required",1),
 ("Emits fumes when overheated — ventilate",1),
 ("Extremely hygroscopic — dry-box storage essential",1)],
 kicker="K2 · MATERIALS HEAD-TO-HEAD",lhead="PETG",rhead="Nylon")
table_slide("Required Temperatures by Material",
 ["Material","Extruder","Bed"],
 [["PLA","195–215 °C","Room temp or 60 °C"],
  ["PETG","240–260 °C","Room temp small parts, or 80–90 °C"],
  ["Nylon","240–260 °C","Room temp, but needs special platforms"],
  ["ABS","240–260 °C","110–120 °C"],
  ["Polycarbonate","280–300 °C","130 °C+"]],
 kicker="K2 · THE WINDOWS",widths=[0.3,0.35,0.35],
 note="Start mid-window; tune in 5 °C steps — strings mean too hot, poor bonding means too cold (Lab 10).")
img_points("Material Settings in Cura",_legacy("s100_da6117bd.png"),[
 ("Material definition","Temperature, fan and flow defaults come from the material profile."),
 ("Settings override","A value set in Print Settings beats the material definition for that print."),
 ("Profiles persist","Create/update a profile to make your tuned overrides permanent."),
 ("One source of truth","Keep tuned profiles per material — the end of guesswork.")],
 kicker="K2 · CONFIGURATION",caption="Cura material management — the definition your prints inherit")
flow_h("Hybrid Metal + Non-Metal Components",[
 "SPLIT — separate the design into metallic and non-metallic parts",
 "SELECT — DMLS/SLM for metals; FDM/SLS for polymers",
 "POST-PROCESS — apply each material's finishing chain separately",
 "BOND — join with the right technique, e.g. electroless deposition",
 "VERIFY — test the assembly's integrity as one component"],
 kicker="A3 · MIXED MATERIALS",color=TEAL)

# ---- Topic 2 labs
cards3("Hands-On Labs — Topic 2",[
 (BLUE,"Toolchain",["Lab 7: Install Cura","Slice your first model","Read the G-code preview"]),
 (TEAL,"Adhesion & support",["Lab 8: Skirt vs brim vs raft","Lab 9: Supports for overhangs"]),
 (VIOLET,"Materials",["Lab 10: Filament diagnostics","Diameter, temperature, flow"])],
 kicker="WHAT YOU'LL DO · LABS 7–10")
for a in DOMAIN2:
    kick=f"TOPIC 02 · HANDS-ON"
    mark(f"lab{a['num']}")
    activity_overview(a,kick)
    lab_visual(a,kick)
    test_slide(a,f"LAB {a['num']} · VERIFY")
content("Recap — Topic 2",[
 "AM issues trace to five sources: hardware, materials, process conditions, environment and design.",
 "The slicer converts STL to G-code — adhesion aids (skirt/brim/raft) and supports are slicing decisions.",
 "Extrusion faults, retraction, oozing, jams and heat creep are diagnosed from the printed evidence.",
 "Each material demands its temperature window, bed and storage — verified, not assumed (Lab 10).",
 "Hybrid metal/polymer components pair the right AM process per part with the right bonding technique."],
 kicker="TOPIC 02 · RECAP",size=16)

# ================= TOPIC 3 =================
T=C.TOPICS[2]
mark("topic3")
section("TOPIC 03",T["title"],"03",T["subtitle"])
img_points("Infill and Perimeters",_legacy("s103_fbe7b8a4.png"),[
 ("Perimeter (walls)","The solid outline around each layer — walls carry most of a part's strength."),
 ("Infill","The internal lattice — typically only 15–20% of the enclosed volume."),
 ("The dial","More walls/infill: stronger, heavier, slower. Less: faster, lighter, weaker."),
 ("Lab 11","You'll measure exactly what each setting costs in time and grams.")],
 kicker="A5 · STRUCTURE",caption="Cura print settings — walls and infill density")
img_points("Infill — Density and Pattern Effects",_legacy("s105_b8ae1465.png"),[
 ("12–20% typical","Fraction of the interior that is plastic — the default band for most parts."),
 ("50%+ for strength","Functional, loaded parts justify denser infill and less warping."),
 ("Pattern matters","Grid, zigzag, gyroid, cubic — each trades strength direction vs speed."),
 ("Lab 12","Zigzag 20% vs grid 30% — printed and squeezed side by side.")],
 kicker="A5 · TUNING",caption="Infill preview at increasing densities")
img_points("Surface Quality — Layer Height",_legacy("s109_6267ac3f.png"),[
 ("The resolution dial","Layer height sets the visible stair-step on every curved or sloped face."),
 ("Time trade","0.1 mm looks superb and takes twice as long as 0.2 mm; 0.3 mm is draft mode."),
 ("Top thickness follows","Keep top THICKNESS constant so the layer count adapts — Lab 13."),
 ("Match to purpose","Draft iterations coarse; presentation parts fine.")],
 kicker="A5 · RESOLUTION",caption="Layer-height setting drives finish and duration")
img_points("The Top Layer — Pillowing",_legacy("s107_1bc595e9.png"),[
 ("Preview lies","The simulated top looks flat; physics prints it over air gaps."),
 ("Pillowing","Thin roofs sag between infill lines, or curl and break — bumpy tops."),
 ("Fixes","More top layers, denser infill under roofs, better cooling."),
 ("Rule of thumb","Top thickness ≥ 4× layer height for a clean roof.")],
 kicker="A5 · ROOFS",caption="Pillowing: the roof sags into sparse infill")
img_points("Surface Texture Options",_legacy("s117_f12b8a65.png"),[
 ("Build surface finish","Glossy or textured bottom faces — decided by the bed surface itself."),
 ("Ironing — Lab 14","The hot nozzle re-passes top faces at ~10% flow for glass-smooth roofs."),
 ("Fuzzy skin — Lab 15","Random jitter on outer walls hides layer lines behind texture."),
 ("Deliberate design","Texture is a slicing decision, not an accident.")],
 kicker="A5 · FINISH",caption="Surface finish elevates a print into a product")
img_points("Support Strategy",_legacy("s116_b78cda9c.png"),[
 ("Overhangs need support","Anything printing into free space beyond ~45–50° needs scaffolding."),
 ("Reorient first","Rotating the part is free — supports cost material, time and finish."),
 ("Placement control","Touching-buildplate vs everywhere; blockers protect chosen faces."),
 ("Removal is a cost","Every support leaves a scar — plan removal before slicing (Lab 16).")],
 kicker="A5 · SUPPORTS",caption="Cura support generation — angle and placement settings")
img_points("Support Challenges",_legacy("s115_975e55ee.png"),[
 ("Inaccessible cavities","Supports inside holes and channels may be impossible to extract."),
 ("Skip or limit","Sometimes the right call is limiting supports and accepting minor sag."),
 ("Design around it","Chamfer the overhang, bridge the gap, or split the part for printing."),
 ("The craft","Support judgement separates operators from process engineers.")],
 kicker="A5 · JUDGEMENT",caption="A part whose internal void must stay support-free")
two_col("Pros and Cons — SLS & Inkjet-Style AM",[
 ("SLS",0),
 ("No supports — powder holds the part",1),
 ("Tough, functional nylon parts; batch nesting",1),
 ("Grainy surface; powder handling & recycling",1),
 ("Long cool-down before parts emerge",1)],
 [("Material/Binder jetting",0),
 ("Multi-material and full-colour capability",1),
 ("Fine features, smooth surfaces",1),
 ("Fragile green parts (binder); resin cost",1),
 ("Nozzle clogging and layer misalignment risks",1)],
 kicker="K5/K6 · TECHNOLOGY REVIEW",lhead="Powder + laser",rhead="Droplets + binder")
two_col("Pros and Cons — DMLS & SLM",[
 ("DMLS — sintered metal",0),
 ("Complex metal geometries, good strength",1),
 ("Established in dental & aerospace",1),
 ("Warping from rapid cooling; powder inconsistency",1),
 ("Expensive machines and powders",1)],
 [("SLM — fully melted metal",0),
 ("Full density, best mechanical properties",1),
 ("Single-alloy purity",1),
 ("Oxidation risk — inert atmosphere required",1),
 ("High residual stresses — heat treatment mandatory",1)],
 kicker="K5/K6 · METAL REVIEW",lhead="Sinter",rhead="Melt")
tile_grid("Business Considerations for AM Operations",[
 ("Production speed","Compare cycle times before and after every process change."),
 ("Quality","Functionality, mechanical properties, dimensional accuracy, surface finish."),
 ("Material usage","Consumption and wastage — supports, purge, failed prints."),
 ("Resource utilisation","Machine uptime, operator hours, floor space, energy."),
 ("Cost & savings","Full cost delta of the change, capital to consumables."),
 ("Ease of use & feedback","Operator experience and the challenges users report.")],
 kicker="K5 · MEASURE THE IMPACT · A5",cols=2,size=13)
flow_h("Assessing a Process Improvement",[
 "BASELINE — record speed, quality, material and cost before",
 "CHANGE — introduce one improvement at a time",
 "MEASURE — same metrics, same parts, after",
 "COMPARE — quantify the delta; collect operator feedback",
 "DECIDE — adopt, tune or roll back; document the standard"],
 kicker="K5 · THE LOOP",color=BLUE)
tile_grid("Sustaining AM Operations Long-Term",[
 ("Calibration & maintenance","Regular calibration and preventive maintenance of AM equipment."),
 ("People","Continuous training and education of AM personnel."),
 ("Materials R&D","Ongoing research into sustainable, environmentally-friendly materials."),
 ("Collaboration","Work with stakeholders on sustainability standards and best practices."),
 ("Circularity","Recycle powders and failed prints; design for material recovery."),
 ("Knowledge base","Document tuned profiles and fixes — institutional memory compounds.")],
 kicker="K7 · SUSTAINABILITY",cols=2,size=13)
tile_grid("Post-Processing — FDM",[
 ("Support removal","Manual removal with tools, or dissolve PVA in water / HIPS in limonene."),
 ("Filling & sanding","Epoxy fillers close gaps; progressive grits smooth surfaces (Lab 16)."),
 ("Vapour smoothing & dipping","Solvent vapour or acetone dip melts the surface to a gloss."),
 ("Priming & painting","Primer then paint for presentation parts."),
 ("Coating","Epoxy resin coats add a protective, sealed skin."),
 ("Metal plating","Chemical or electro-plating gives polymer parts a metallic finish.")],
 kicker="A5 · FINISHING FDM",cols=2,size=13)
tile_grid("Post-Processing — SLS and SLA",[
 ("SLS: cooling","Hours of controlled cool-down before parts leave the build chamber."),
 ("SLS: clean-up","Compressed air, vacuum, sand blasting or brushing recover unsintered powder."),
 ("SLS: finishing","Sanding, polishing, dyeing, painting, coatings, electroplating, vapour smoothing."),
 ("SLA: wash & cure","IPA wash removes resin; UV post-cure completes polymerisation."),
 ("SLA: finishing","Sanding, polishing, dyeing, painting, coating, electroplating."),
 ("Safety","Gloves and ventilation — powders and uncured resin are hazardous.")],
 kicker="A5 · FINISHING POWDER & RESIN",cols=2,size=13)
tile_grid("Post-Processing — Metal AM",[
 ("Sieving & depowdering","Unmelted powder recovered by sieving; trapped powder shaken out by vibration."),
 ("Heat treatment","Vacuum/inert furnace relieves thermal stress — before build-plate separation."),
 ("Part cut-off","Chiselling, bandsaw or wire EDM separate the part from the plate."),
 ("Support removal","Chiselling, wire EDM, CNC machining or thermal deburring by geometry."),
 ("Machining","Threads, holes and precision flats machined to tolerance."),
 ("Surface finish","Shot blasting, plating, polishing, micromachining — per the specification.")],
 kicker="A5 · FINISHING METALS",cols=2,size=13)

# ---- Topic 3 labs
cards3("Hands-On Labs — Topic 3",[
 (BLUE,"Strength & speed",["Lab 11: Perimeters & infill","Lab 12: Infill patterns strength test"]),
 (TEAL,"Surface quality",["Lab 13: Layer height ladder","Lab 14: Ironing","Lab 15: Fuzzy skin"]),
 (VIOLET,"Finishing",["Lab 16: Support removal","Progressive sanding to smooth"])],
 kicker="WHAT YOU'LL DO · LABS 11–16")
for a in DOMAIN3:
    kick=f"TOPIC 03 · HANDS-ON"
    mark(f"lab{a['num']}")
    activity_overview(a,kick)
    if a["num"] not in (11,13):   # their reference images already appear on the concept slides
        lab_visual(a,kick)
    test_slide(a,f"LAB {a['num']} · VERIFY")
content("Recap — Topic 3",[
 "Walls and infill are the strength-vs-speed dial; measure what every setting costs before standardising it.",
 "Layer height, top thickness, ironing and fuzzy skin put surface quality under deliberate control.",
 "Reorient first, support second — and plan support removal before slicing, not after printing.",
 "Assess every improvement with the baseline→change→measure→compare→decide loop across speed, quality, material, cost and feedback.",
 "Calibration, training, sustainable materials and collaboration keep an AM operation viable long-term."],
 kicker="TOPIC 03 · RECAP",size=16)

# ---------------- CLOSE ----------------
mark("close")
section("WRAP-UP","Course Summary & Next Steps","")
tile_grid("What You Achieved",[
 ("Analysed AM systems","Seven process families, materials, benefits and trade-offs — and when to choose which."),
 ("Designed for AM","Extrude, sweep, loft, revolve, tolerances and stress-aware design in Fusion 360."),
 ("Reverse engineered","3D-scanned a real object and rebuilt scans into printable CAD."),
 ("Diagnosed process issues","Adhesion, extrusion, temperature, jams and materials — evidence-based fixes."),
 ("Optimised processes","Infill, surface quality, supports and post-processing — measured, not guessed."),
 ("Assessed business impact","Speed, quality, material, cost and feedback before and after each change.")],
 kicker="LEARNING OUTCOMES DELIVERED",cols=2,size=13)
img_points("Practice Exam & Revision",_legacy("s015_8266d9af.png"),[
 ("Practice portal","Sharpen your readiness at https://exams.tertiaryinfotech.com."),
 ("Revise the labs","Redo each lab's Test-it check from memory — settings and reasons."),
 ("Know the frameworks","7 process families · 5 issue sources · the improvement loop."),
 ("Open book ≠ open mind","Bookmark your slides and Learner Guide sections before the assessment.")],
 kicker="TEST YOURSELF",caption="exams.tertiaryinfotech.com · lms-tms.tertiaryinfotech.com")
img_points("Certificate & TRAQOM Survey",_legacy("s142_59bff3df.png"),[
 ("Certificate of Achievement","Issued on 75% attendance + Competent assessment — download from the LMS."),
 ("OpenCert","SkillsFuture Singapore issues an OpenCert Statement of Attainment."),
 ("TRAQOM survey","Scan the survey QR — your feedback is mandatory for WSQ funding."),
 ("Portal","https://lms-tms.tertiaryinfotech.com")],
 kicker="WRAP-UP · MANDATORY",caption="Certificate delivery and TRAQOM survey QR codes on the LMS")
content("Recommended Next Courses",[
 "WSQ — Product Design with Fusion 360",
 "WSQ — Product Design with Autodesk Inventor",
 "WSQ — AutoCAD Civil 3D for Infrastructure Design",
 "WSQ — Basic Technical Drawing with SketchUp for Beginners",
 "WSQ — 3D Modelling with Blender for Beginners"],kicker="KEEP BUILDING")
content("Support",[
 "If you have any enquiries during and after the class, contact us:",
 "Email: enquiry@tertiaryinfotech.com",
 "Tel: 6100 0613",
 "Web: www.tertiarycourses.com.sg"],kicker="WE'RE HERE TO HELP")
tile_grid("Assessment",[
 ("Written Assessment (WA)","Short-Answer Questions (SAQ) — 1 hour, open book."),
 ("Practical Performance (PP)","Scenario-based practical tasks — 90 minutes, open book."),
 ("Open book","Slides, Learner Guide and approved materials only."),
 ("Digital attendance","Remember to take the Assessment digital attendance (TRAQOM)."),
 ("Eligibility","Minimum 75% attendance to be eligible for assessment and funding."),
 ("Upload","Submit your completed answers on https://lms-tms.tertiaryinfotech.com.")],
 kicker="FINAL ASSESSMENT",cols=2,size=13)
flow_h("Assessment Flow",[
 "TRAQOM survey — scan the QR code on the LMS",
 "Assessment digital attendance — scan the SSG QR",
 "Sit WA (SAQ) then the Practical Performance — open book",
 "Upload your answers on the LMS",
 "Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY")
tile_grid("Digital Attendance (Mandatory)",[
 ("AM · PM · Assessment","It is mandatory to take the AM, PM and Assessment digital attendance for WSQ-funded courses."),
 ("SSG QR code","The trainer/administrator displays the digital attendance QR code generated from the SSG portal."),
 ("Scan & submit","Scan the QR code with your mobile phone camera and submit your attendance."),
 ("75% minimum","A minimum of 75% attendance is required to be eligible for assessment and funding.")],
 kicker="TRAQOM · SSG DIGITAL ATTENDANCE",cols=2,size=15)
big_statement("Thank You!","You are now ready to innovate product development with 3D printing and additive manufacturing.","HAPPY PRINTING",color=TEAL)

# ---------------- transitions (one pass) ----------------
def _transition(s,kind="fade",speed="fast"):
    sld=s._element
    for old in sld.findall(qn("p:transition")): sld.remove(old)
    tr=etree.SubElement(sld,qn("p:transition"))
    tr.set("spd",speed); tr.set("advClick","1")
    if kind=="fade": etree.SubElement(tr,qn("p:fade"))
    elif kind=="push": etree.SubElement(tr,qn("p:push")).set("dir","l")
    sld.append(tr)
SECTION_TITLES=set()
for i,s in enumerate(prs.slides):
    is_divider=any(sh.has_text_frame and sh.text_frame.text.strip().startswith(("TOPIC 0","COURSE ADMINISTRATION","WRAP-UP"))
                   for sh in s.shapes if sh.has_text_frame)
    _transition(s,"push" if is_divider else "fade","med" if is_divider else "fast")

OUT=os.path.join(REPO,"courseware",f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
prs.save(OUT)
SLIDE_MAP["total"]=len(prs.slides._sldIdLst)
with open(os.path.join(HERE,"slide_map.json"),"w") as f:
    json.dump(SLIDE_MAP,f,indent=1)
print(f"Saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")
